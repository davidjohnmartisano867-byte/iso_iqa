import os
import io
import json
import re
import tempfile
from typing import Optional, List, Any
from fastapi import FastAPI, File, UploadFile, Form, Body, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

app = FastAPI(title="ISO-QA AI Microservice", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

CLAUSE_TITLES = {
    1: "Scope",
    2: "Normative References",
    3: "Terms and Definitions",
    4: "Context of the Organization",
    5: "Leadership",
    6: "Planning",
    7: "Support",
    8: "Operation",
    9: "Performance Evaluation",
    10: "Improvement"
}

def extract_text_from_bytes(filename: str, content: bytes) -> str:
    ext = os.path.splitext(filename)[1].lower()
    text = ""
    try:
        if ext == '.pdf':
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text += t + " "
        elif ext in ['.doc', '.docx']:
            import docx
            doc = docx.Document(io.BytesIO(content))
            for p in doc.paragraphs:
                if p.text:
                    text += p.text + " "
        elif ext in ['.xls', '.xlsx']:
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    for cell in row:
                        if cell is not None:
                            text += str(cell) + " "
        else:
            text = content.decode('utf-8', errors='ignore')
    except Exception:
        pass
        
    if not text.strip():
        text = filename
    return text

def keyword_classify(text: str):
    keywords = {
        1: ["scope", "applicability", "boundaries", "requirements"],
        2: ["normative", "references", "standards", "iso"],
        3: ["terms", "definitions", "abbreviations", "glossary"],
        4: ["context", "organization", "stakeholder", "swot", "pestle", "interested parties"],
        5: ["leadership", "policy", "commitment", "officer", "manager", "role", "responsibility", "authority"],
        6: ["planning", "objective", "risk", "mitigation", "opportunity", "action", "change"],
        7: ["support", "resource", "competence", "infrastructure", "awareness", "communication", "document", "information"],
        8: ["operation", "process", "design", "production", "service", "execution", "control", "release", "output"],
        9: ["performance", "evaluation", "monitor", "measure", "analysis", "audit", "satisfaction", "review", "metrics"],
        10: ["improvement", "nonconformity", "corrective", "action", "continual", "innovation"]
    }
    scores = {i: 0 for i in range(1, 11)}
    text_lower = text.lower()
    for clause, words in keywords.items():
        for word in words:
            scores[clause] += len(re.findall(r'\b' + re.escape(word) + r'\b', text_lower))
    best = max(scores, key=scores.get)
    if scores[best] == 0:
        return 8, "Defaulted to Operations. No clear keywords matched."
    return best, f"Classified based on keyword matching patterns (Top score: {scores[best]} hits)."

def check_document_flags(text: str, filename: str):
    text_lower = text.lower()
    fname_lower = filename.lower()
    is_fake = False
    lacks_signature = False
    unrelated_to_office = False

    fake_terms = ['lorem ipsum', 'consectetur adipiscing', 'sample text', 'placeholder text', 
                  'test test', 'dummy document', 'fake file', 'temp file', 'asdf', 'qwerty', 'foo bar']
    if any(term in text_lower for term in fake_terms) or len(text_lower.strip()) < 10:
        is_fake = True

    formal_markers = ['memo', 'memorandum', 'policy', 'procedure', 'form', 'report', 
                       'request', 'plan', 'manual', 'audit', 'iso', 'certificate', 'agreement']
    signature_markers = ['signature', 'approved by', 'signed by', 'recommending approval', 
                         'officer-in-charge', 'approved:', 'prepared by:', 'sign-off', '/s/', 'signed']
    is_formal = any(m in fname_lower or m in text_lower for m in formal_markers)
    has_sig = any(s in text_lower for s in signature_markers)
    if is_formal and not has_sig:
        lacks_signature = True

    work_keywords = ['office', 'department', 'iso', 'quality', 'audit', 'process', 'management', 
                     'policy', 'procedure', 'report', 'academic', 'student', 'faculty', 'university', 
                     'administration', 'system', 'data', 'record', 'file', 'service', 'operation', 
                     'log', 'form', 'memo', 'request', 'schedule', 'compliance', 'clause', 'finding', 
                     'capa', 'auditor', 'auditee', 'hr', 'finance', 'document']
    work_hits = sum(1 for kw in work_keywords if kw in text_lower)
    if work_hits < 1 and not is_formal:
        unrelated_to_office = True

    warnings = []
    if is_fake:
        warnings.append("⚠️ Document appears to contain fake, dummy, or placeholder content.")
    if lacks_signature:
        warnings.append("⚠️ Missing Signature: Document lacks an official signature or approval sign-off block.")
    if unrelated_to_office:
        warnings.append("⚠️ Irrelevant Content: Document does not appear to be related to official office or work operations.")

    if warnings:
        comment = "AI Quality Verification Notice:\n" + "\n".join(warnings)
    else:
        comment = "✅ AI Quality Verification: Document verified as official office documentation."

    return {
        "is_fake": is_fake,
        "lacks_signature": lacks_signature,
        "unrelated_to_office": unrelated_to_office,
        "ai_comment": comment
    }

def get_gemini_api_key():
    key = os.environ.get("GEMINI_API_KEY")
    if not key:
        env_candidates = [
            os.path.join(os.path.dirname(__file__), ".env"),
            os.path.join(os.path.dirname(__file__), "..", ".env"),
        ]
        for env_path in env_candidates:
            if os.path.exists(env_path):
                with open(env_path, "r", encoding="utf-8") as f:
                    for line in f:
                        if line.startswith("GEMINI_API_KEY="):
                            return line.split("=", 1)[1].strip()
    return key

@app.get("/")
def root():
    return {
        "status": "online",
        "service": "ISO-QA AI Microservice",
        "endpoints": {
            "/classify": "POST document file or text for AI classification & verification",
            "/predict": "POST office historical compliance data for predictive risk & clause forecasting",
            "/health": "GET health status"
        }
    }

@app.get("/health")
def health():
    return {"status": "ok"}

@app.post("/classify")
async def classify_document_endpoint(
    file: Optional[UploadFile] = File(None),
    filename: Optional[str] = Form(None),
    text_content: Optional[str] = Form(None),
    clause: Optional[str] = Form("auto")
):
    actual_filename = filename or (file.filename if file else "document.txt")
    if file:
        content_bytes = await file.read()
        text = extract_text_from_bytes(actual_filename, content_bytes)
    elif text_content:
        text = text_content
    else:
        raise HTTPException(status_code=400, detail="No file or text provided")

    fallback_flags = check_document_flags(text, actual_filename)
    is_auto = (clause == "auto")
    api_key = get_gemini_api_key()

    if not api_key:
        c_num, reasoning = keyword_classify(text)
        if not is_auto and clause and clause.isdigit():
            c_num = int(clause)
            reasoning = f"Manually assigned to Clause {c_num}."
        return {
            "success": True,
            "clause_number": c_num,
            "clause_title": CLAUSE_TITLES.get(c_num, "Operation"),
            "reasoning": reasoning + " (API key not configured)",
            "ai_comment": fallback_flags["ai_comment"],
            "is_fake": fallback_flags["is_fake"],
            "lacks_signature": fallback_flags["lacks_signature"],
            "unrelated_to_office": fallback_flags["unrelated_to_office"]
        }

    try:
        import google.generativeai as genai
        genai.configure(api_key=api_key)

        prompt = f"""You are an ISO 9001:2015 Quality Management AI Auditor.
Analyze this uploaded document for an organization repository.

Tasks:
1. Classify under one of the 10 ISO 9001:2015 Clauses (1 to 10).
2. Check for quality issues:
   - is_fake: dummy text, lorem ipsum, fake, corrupt, or placeholder data?
   - lacks_signature: formal document/policy/report that lacks an official signature or approval line?
   - unrelated_to_office: unrelated to work, university, or office operations?
3. Provide a clear comment (ai_comment) summarizing your findings.

File Name: {actual_filename}
Content:
{text[:4000]}

Respond with a valid JSON object ONLY:
{{
  "clause_number": 8,
  "reasoning": "Short explanation",
  "is_fake": false,
  "lacks_signature": false,
  "unrelated_to_office": false,
  "ai_comment": "Clear verification note"
}}
"""
        model = genai.GenerativeModel("gemini-2.5-flash")
        resp = model.generate_content(prompt)
        resp_text = resp.text.strip()
        if resp_text.startswith("```json"):
            resp_text = resp_text[7:]
        if resp_text.startswith("```"):
            resp_text = resp_text[3:]
        if resp_text.endswith("```"):
            resp_text = resp_text[:-3]
        data = json.loads(resp_text.strip())

        c_num = int(data.get("clause_number", 8))
        if not is_auto and clause and str(clause).isdigit():
            c_num = int(clause)

        return {
            "success": True,
            "clause_number": c_num,
            "clause_title": CLAUSE_TITLES.get(c_num, "Operation"),
            "reasoning": data.get("reasoning", "Classified with Gemini AI"),
            "ai_comment": data.get("ai_comment") or fallback_flags["ai_comment"],
            "is_fake": bool(data.get("is_fake", fallback_flags["is_fake"])),
            "lacks_signature": bool(data.get("lacks_signature", fallback_flags["lacks_signature"])),
            "unrelated_to_office": bool(data.get("unrelated_to_office", fallback_flags["unrelated_to_office"]))
        }
    except Exception as e:
        c_num, reasoning = keyword_classify(text)
        if not is_auto and clause and str(clause).isdigit():
            c_num = int(clause)
        return {
            "success": True,
            "clause_number": c_num,
            "clause_title": CLAUSE_TITLES.get(c_num, "Operation"),
            "reasoning": f"Fallback: {str(e)}",
            "ai_comment": fallback_flags["ai_comment"],
            "is_fake": fallback_flags["is_fake"],
            "lacks_signature": fallback_flags["lacks_signature"],
            "unrelated_to_office": fallback_flags["unrelated_to_office"]
        }

@app.post("/predict")
async def predict_endpoint(payload: List[Any] = Body(...)):
    try:
        try:
            from predict_engine import predict_with_scikit_learn, generate_hybrid_forecast
        except ImportError:
            from ai_service.predict_engine import predict_with_scikit_learn, generate_hybrid_forecast

        ml_preds = predict_with_scikit_learn(payload)
        final_results = generate_hybrid_forecast(payload, ml_preds)
        return {
            "success": True,
            "predictions": final_results
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Prediction error: {str(e)}")

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="0.0.0.0", port=8000, reload=True)
