import os
import sys
import json
import re

def extract_text(file_path):
    ext = os.path.splitext(file_path)[1].lower()
    text = ""
    try:
        if ext == '.pdf':
            import pypdf
            reader = pypdf.PdfReader(file_path)
            for page in reader.pages:
                t = page.extract_text()
                if t:
                    text += t + " "
        elif ext in ['.doc', '.docx']:
            import docx
            doc = docx.Document(file_path)
            for p in doc.paragraphs:
                if p.text:
                    text += p.text + " "
        elif ext in ['.xls', '.xlsx']:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            for sheet in wb.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    for cell in row:
                        if cell is not None:
                            text += str(cell) + " "
        else:
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                text = f.read()
    except Exception as e:
        pass
    
    if not text.strip():
        text = os.path.basename(file_path)
    return text

def keyword_classify(text):
    keywords = {
        1: ["scope", "applicability", "boundaries", "requirements"],
        2: ["normative", "references", "standards", "iso"],
        3: ["terms", "definitions", "abbreviations", "glossary"],
        4: ["context", "organization", "stakeholder", "swot", "pestle", "interested parties"],
        5: ["leadership", "policy", "commitment", "officer", "manager", "role", "responsibility", "authority"],
        6: ["planning", "objective", "risk", "mitigation", "opportunity", "action", "change"],
        7: ["support", "resource", "competence", "infrastructure", "awareness", "communication", "document", "information"],
        8: ["operation", "process", "design", "production", "service", "execution", "control", "release", "output"],
        9: ["performance", "evaluation", "monitor", "measure", "analysis", "audit", "satisfaction", "review", "metrics"],
        10: ["improvement", "nonconformity", "corrective", "action", "continual", "innovation"]
    }
    
    scores = {i: 0 for i in range(1, 11)}
    text_lower = text.lower()
    for clause, words in keywords.items():
        for word in words:
            scores[clause] += len(re.findall(r'\b' + re.escape(word) + r'\b', text_lower))
            
    best_clause = max(scores, key=scores.get)
    if scores[best_clause] == 0:
        return 8, "Defaulted to Operations. No clear keywords matched."
    return best_clause, f"Classified based on keyword matching patterns (Top score: {scores[best_clause]} hits)."

def check_document_flags(text, filename):
    text_lower = text.lower()
    fname_lower = filename.lower()
    
    is_fake = False
    lacks_signature = False
    unrelated_to_office = False
    
    fake_terms = ['lorem ipsum', 'consectetur adipiscing', 'sample text', 'placeholder text', 
                  'test test', 'dummy document', 'fake file', 'temp file', 'asdf', 'qwerty', 'foo bar']
    if any(term in text_lower for term in fake_terms) or len(text_lower.strip()) < 10:
        is_fake = True
        
    formal_markers = ['memo', 'memorandum', 'policy', 'procedure', 'form', 'report', 
                       'request', 'plan', 'manual', 'audit', 'iso', 'certificate', 'agreement']
    signature_markers = ['signature', 'approved by', 'signed by', 'recommending approval', 
                         'officer-in-charge', 'approved:', 'prepared by:', 'sign-off', '/s/', 'signed']
    
    is_formal = any(m in fname_lower or m in text_lower for m in formal_markers)
    has_sig = any(s in text_lower for s in signature_markers)
    if is_formal and not has_sig:
        lacks_signature = True
        
    work_keywords = ['office', 'department', 'iso', 'quality', 'audit', 'process', 'management', 
                     'policy', 'procedure', 'report', 'academic', 'student', 'faculty', 'university', 
                     'administration', 'system', 'data', 'record', 'file', 'service', 'operation', 
                     'log', 'form', 'memo', 'request', 'schedule', 'compliance', 'clause', 'finding', 
                     'capa', 'auditor', 'auditee', 'hr', 'finance', 'document']
    work_hits = sum(1 for kw in work_keywords if kw in text_lower)
    if work_hits < 1 and not is_formal:
        unrelated_to_office = True
        
    warnings = []
    if is_fake:
        warnings.append("⚠️ Document appears to contain fake, dummy, or placeholder content.")
    if lacks_signature:
        warnings.append("⚠️ Missing Signature: Document lacks an official signature or approval sign-off block.")
    if unrelated_to_office:
        warnings.append("⚠️ Irrelevant Content: Document does not appear to be related to official office or work operations.")
        
    if warnings:
        comment = "AI Quality Verification Notice:\n" + "\n".join(warnings)
    else:
        comment = "✅ AI Quality Verification: Document verified as official office documentation."
        
    return {
        "is_fake": is_fake,
        "lacks_signature": lacks_signature,
        "unrelated_to_office": unrelated_to_office,
        "ai_comment": comment
    }

def main():
    if len(sys.argv) < 2:
        print(json.dumps({"clause_number": 8, "reasoning": "No file path provided.", "ai_comment": "No file path provided."}))
        sys.exit(1)
        
    file_path = sys.argv[1]
    if not os.path.exists(file_path):
        print(json.dumps({"clause_number": 8, "reasoning": "File does not exist.", "ai_comment": "File does not exist."}))
        sys.exit(1)
        
    text = extract_text(file_path)
    filename = os.path.basename(file_path)
    fallback_flags = check_document_flags(text, filename)
    
    api_key = os.environ.get("GEMINI_API_KEY")
    
    if not api_key:
        env_path = os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env")
        if os.path.exists(env_path):
            with open(env_path, "r") as f:
                for line in f:
                    if line.startswith("GEMINI_API_KEY="):
                        api_key = line.split("=", 1)[1].strip()
                        break
                        
    if not api_key:
        clause, reasoning = keyword_classify(text)
        result = {
            "clause_number": clause,
            "reasoning": reasoning + " (API key not configured)",
            "ai_comment": fallback_flags["ai_comment"],
            "is_fake": fallback_flags["is_fake"],
            "lacks_signature": fallback_flags["lacks_signature"],
            "unrelated_to_office": fallback_flags["unrelated_to_office"]
        }
        print(json.dumps(result))
        sys.exit(0)
        
    try:
        import google.generativeai as genai
        genai.configure(api_key=api_key)
        
        prompt = """You are an ISO 9001:2015 Quality Management AI Auditor.
Your task is to analyze the content of the uploaded document for an office/organization repository.

Tasks:
1. Classify the document under one of the 10 ISO 9001:2015 Clauses (1 to 10).
2. Check for quality issues:
   - is_fake: Is the content dummy text, lorem ipsum, fake, corrupt, or placeholder data?
   - lacks_signature: Is it a formal document/policy/report/form that lacks an official signature or approval line?
   - unrelated_to_office: Is the content unrelated to work, university operations, or office documentation (e.g. personal notes, recipes, random media)?
3. Provide a clear, professional comment (ai_comment) summarizing your findings or warnings.

Analyze the following document metadata and content excerpt:
File Name: __FILENAME__
Content:
__CONTENT__

Respond with a valid JSON object ONLY. Do not include markdown formatting or backticks.
JSON Structure:
{
  "clause_number": 8,
  "reasoning": "Short one-sentence explanation of why it fits this clause",
  "is_fake": false,
  "lacks_signature": true,
  "unrelated_to_office": false,
  "ai_comment": "AI Verification Comment summarizing any warnings (or positive verification)."
}
"""
        prompt = prompt.replace("__FILENAME__", filename).replace("__CONTENT__", text[:4000])
        
        response = None
        last_err = None
        for model_name in ["gemini-1.5-flash", "gemini-2.5-flash", "gemini-pro"]:
            try:
                model = genai.GenerativeModel(model_name)
                response = model.generate_content(prompt)
                break
            except Exception as e:
                last_err = e
                
        if not response:
            raise last_err
            
        response_text = response.text.strip()
        
        if "```json" in response_text:
            response_text = response_text.split("```json")[1].split("```")[0].strip()
        elif "```" in response_text:
            response_text = response_text.split("```")[1].split("```")[0].strip()
            
        data = json.loads(response_text)
        if "ai_comment" not in data or not data["ai_comment"]:
            data["ai_comment"] = fallback_flags["ai_comment"]
        print(json.dumps(data))
    except Exception as e:
        clause, reasoning = keyword_classify(text)
        result = {
            "clause_number": clause,
            "reasoning": f"Fallback due to Gemini API error: {str(e)}",
            "ai_comment": fallback_flags["ai_comment"],
            "is_fake": fallback_flags["is_fake"],
            "lacks_signature": fallback_flags["lacks_signature"],
            "unrelated_to_office": fallback_flags["unrelated_to_office"]
        }
        print(json.dumps(result))

if __name__ == "__main__":
    main()
